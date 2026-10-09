from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from omni_analyst.ingestion.models import (
    DocumentSource,
    LayoutBlock,
    PageLayout,
    ParsedDocument,
    ParsedElement,
    TableCell,
    TableRecord,
)
from omni_analyst.ingestion.parsers.base import BaseParser


class XlsxParser(BaseParser):
    parser_name = "xlsx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    extensions = (".xlsx",)

    def parse(self, source: DocumentSource) -> ParsedDocument:
        workbook = load_workbook(
            filename=BytesIO(source.read_bytes()),
            read_only=True,
            data_only=True,
        )
        elements: list[ParsedElement] = []
        tables: list[TableRecord] = []
        pages: list[PageLayout] = []
        row_lines: list[str] = []

        try:
            for worksheet in workbook.worksheets:
                table_rows: list[list[TableCell]] = []
                sheet_lines: list[str] = []
                for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                    values = [self._format_cell(value) for value in row]
                    line = " | ".join(value for value in values if value)
                    if not line:
                        continue
                    row_lines.append(line)
                    sheet_lines.append(line)
                    table_rows.append(
                        [
                            TableCell(
                                row_index=row_index,
                                column_index=column_index,
                                text=value,
                                metadata={"sheet": worksheet.title, "source": "xlsx_cell"},
                            )
                            for column_index, value in enumerate(values)
                        ]
                    )
                    elements.append(
                        ParsedElement(
                            element_type="sheet_row",
                            text=line,
                            metadata={
                                "sheet": worksheet.title,
                                "row_index": row_index,
                            },
                        )
                    )
                if table_rows:
                    table = TableRecord(
                        table_id=f"{source.document_id}-{worksheet.title}-table",
                        rows=table_rows,
                        source="xlsx",
                        metadata={"sheet": worksheet.title},
                    )
                    tables.append(table)
                    pages.append(
                        PageLayout(
                            page_number=len(pages) + 1,
                            text="\n".join(sheet_lines),
                            blocks=[
                                LayoutBlock(
                                    block_id=f"{source.document_id}-{worksheet.title}-row-{index}",
                                    block_type="sheet_row",
                                    text=line,
                                    page_number=len(pages) + 1,
                                    metadata={"sheet": worksheet.title, "row_index": index},
                                )
                                for index, line in enumerate(sheet_lines)
                            ],
                            tables=[table],
                            metadata={"sheet": worksheet.title},
                        )
                    )
        finally:
            workbook.close()

        return ParsedDocument(
            document_id=source.document_id,
            filename=source.filename,
            media_type=source.media_type
            or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            parser_name=self.parser_name,
            text="\n".join(row_lines),
            elements=elements,
            pages=pages,
            tables=tables,
            metadata={"rows": len(row_lines), "sheets": len(workbook.sheetnames), "table_records": len(tables)},
        )

    def _format_cell(self, value: object) -> str:
        return "" if value is None else str(value)
