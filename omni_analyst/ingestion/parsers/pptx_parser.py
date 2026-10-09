from __future__ import annotations

from io import BytesIO

from pptx import Presentation

from omni_analyst.ingestion.models import (
    DocumentSource,
    LayoutBlock,
    PageLayout,
    ParsedDocument,
    ParsedElement,
    TableCell,
    TableRecord,
)
from omni_analyst.ingestion.parsers.base import BaseParser


class PptxParser(BaseParser):
    parser_name = "pptx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    extensions = (".pptx",)

    def parse(self, source: DocumentSource) -> ParsedDocument:
        presentation = Presentation(BytesIO(source.read_bytes()))
        slide_lines: list[str] = []
        elements: list[ParsedElement] = []
        pages: list[PageLayout] = []
        tables: list[TableRecord] = []

        for slide_index, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            slide_tables: list[TableRecord] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    rows: list[list[TableCell]] = []
                    for row_index, row in enumerate(shape.table.rows):
                        cells = [
                            TableCell(
                                row_index=row_index,
                                column_index=column_index,
                                text=cell.text.strip(),
                                metadata={"slide_index": slide_index, "source": "pptx_table"},
                            )
                            for column_index, cell in enumerate(row.cells)
                        ]
                        rows.append(cells)
                        row_text = " | ".join(cell.text for cell in cells)
                        if row_text.strip():
                            texts.append(row_text)
                    if rows:
                        table = TableRecord(
                            table_id=f"{source.document_id}-slide-{slide_index}-table-{len(slide_tables)}",
                            rows=rows,
                            page_number=slide_index,
                            source="pptx",
                        )
                        slide_tables.append(table)
                        tables.append(table)
            if not texts:
                continue
            slide_text = "\n".join(text for text in texts if text)
            slide_lines.append(f"Slide {slide_index}\n{slide_text}")
            elements.append(
                ParsedElement(
                    element_type="slide",
                    text=slide_text,
                    metadata={"slide_index": slide_index},
                )
            )
            pages.append(
                PageLayout(
                    page_number=slide_index,
                    text=slide_text,
                    blocks=[
                        LayoutBlock(
                            block_id=f"{source.document_id}-slide-{slide_index}-block-{index}",
                            block_type="slide_text",
                            text=text,
                            page_number=slide_index,
                            metadata={"slide_index": slide_index},
                        )
                        for index, text in enumerate(texts)
                    ],
                    tables=slide_tables,
                    metadata={"slide_index": slide_index},
                )
            )

        return ParsedDocument(
            document_id=source.document_id,
            filename=source.filename,
            media_type=source.media_type
            or "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            parser_name=self.parser_name,
            text="\n\n".join(slide_lines),
            elements=elements,
            pages=pages,
            tables=tables,
            metadata={"slides": len(elements), "tables": len(tables)},
        )
