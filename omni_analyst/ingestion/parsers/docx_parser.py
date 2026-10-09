from __future__ import annotations

from io import BytesIO

from docx import Document

from omni_analyst.ingestion.models import (
    DocumentSource,
    LayoutBlock,
    PageLayout,
    ParsedDocument,
    ParsedElement,
    TableCell,
    TableRecord,
)
from omni_analyst.ingestion.parsers.base import BaseParser


class DocxParser(BaseParser):
    parser_name = "docx"
    media_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    extensions = (".docx",)

    def parse(self, source: DocumentSource) -> ParsedDocument:
        document = Document(BytesIO(source.read_bytes()))
        elements: list[ParsedElement] = []
        tables: list[TableRecord] = []

        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            elements.append(
                ParsedElement(
                    element_type="paragraph",
                    text=text,
                    metadata={
                        "paragraph_index": index,
                        "style": paragraph.style.name if paragraph.style else None,
                    },
                )
            )

        for table_index, table in enumerate(document.tables):
            table_rows: list[list[TableCell]] = []
            for row_index, row in enumerate(table.rows):
                cells = [
                    TableCell(
                        row_index=row_index,
                        column_index=column_index,
                        text=cell.text.strip(),
                        metadata={"source": "docx_table"},
                    )
                    for column_index, cell in enumerate(row.cells)
                ]
                table_rows.append(cells)
                row_text = " | ".join(cell.text for cell in cells)
                if not row_text.strip():
                    continue
                elements.append(
                    ParsedElement(
                        element_type="table_row",
                        text=row_text,
                        metadata={
                            "table_index": table_index,
                            "row_index": row_index,
                        },
                    )
                )
            if table_rows:
                tables.append(
                    TableRecord(
                        table_id=f"{source.document_id}-table-{table_index}",
                        rows=table_rows,
                        source="docx",
                    )
                )

        joined = "\n\n".join(element.text for element in elements)
        page = PageLayout(
            page_number=1,
            text=joined,
            blocks=[
                LayoutBlock(
                    block_id=f"{source.document_id}-element-{index}",
                    block_type=element.element_type,
                    text=element.text,
                    page_number=1,
                    metadata=element.metadata,
                )
                for index, element in enumerate(elements)
            ],
            tables=tables,
        )
        return ParsedDocument(
            document_id=source.document_id,
            filename=source.filename,
            media_type=source.media_type
            or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            parser_name=self.parser_name,
            text=joined,
            elements=elements,
            pages=[page],
            tables=tables,
            metadata={
                "paragraphs": sum(1 for element in elements if element.element_type == "paragraph"),
                "tables": len(document.tables),
                "table_records": len(tables),
            },
        )
