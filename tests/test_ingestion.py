from __future__ import annotations

import base64
import io
import unittest

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from omni_analyst.graph.simple_builder import SimpleGraphBuilder
from omni_analyst.ingestion.chunking import AntiChunkingEngine, ChunkQualityMiddleware
from omni_analyst.ingestion.models import DocumentSource, ParsedDocument
from omni_analyst.ingestion.parsers.csv_parser import CsvParser
from omni_analyst.ingestion.parsers.docx_parser import DocxParser
from omni_analyst.ingestion.parsers.pdf_parser import PdfParser
from omni_analyst.ingestion.parsers.pptx_parser import PptxParser
from omni_analyst.ingestion.parsers.xlsx_parser import XlsxParser
from omni_analyst.ingestion.pipeline import IngestionPipeline


class ParserTests(unittest.TestCase):
    def test_csv_parser_extracts_rows(self) -> None:
        parser = CsvParser()
        document = parser.parse(
            DocumentSource(
                document_id="doc-1",
                filename="metrics.csv",
                media_type="text/csv",
                text="name,value\nrevenue,42",
            )
        )
        self.assertIn("revenue | 42", document.text)
        self.assertEqual(document.metadata["rows"], 2)
        self.assertEqual(document.metadata["table_records"], 1)
        self.assertEqual(document.tables[0].rows[1][0].text, "revenue")

    def test_pdf_parser_extracts_native_text_with_pymupdf(self) -> None:
        parser = PdfParser()
        source = DocumentSource(
            document_id="doc-pdf-1",
            filename="native.pdf",
            content_base64=self._b64(self._pdf_bytes("Native revenue text")),
        )
        document = parser.parse(source)
        self.assertIn("Native revenue text", document.text)
        self.assertTrue(document.metadata["has_text_layer"])
        self.assertGreaterEqual(document.metadata["layout_blocks"], 1)
        self.assertEqual(document.pages[0].page_number, 1)

    def test_docx_parser_extracts_paragraphs_and_tables(self) -> None:
        parser = DocxParser()
        source = DocumentSource(
            document_id="doc-2",
            filename="notes.docx",
            content_base64=self._b64(self._docx_bytes()),
        )
        document = parser.parse(source)
        self.assertIn("Hello world", document.text)
        self.assertIn("Vendor | Amount", document.text)
        self.assertEqual(document.metadata["paragraphs"], 2)
        self.assertEqual(document.metadata["tables"], 1)
        self.assertEqual(document.tables[0].rows[1][0].text, "ACME")

    def test_xlsx_parser_extracts_cells(self) -> None:
        parser = XlsxParser()
        source = DocumentSource(
            document_id="doc-3",
            filename="sheet.xlsx",
            content_base64=self._b64(self._xlsx_bytes()),
        )
        document = parser.parse(source)
        self.assertIn("Quarter | Revenue", document.text)
        self.assertIn("Q1 | 120", document.text)
        self.assertEqual(document.metadata["sheets"], 1)
        self.assertEqual(document.tables[0].metadata["sheet"], "Sheet")

    def test_pptx_parser_extracts_slide_text(self) -> None:
        parser = PptxParser()
        source = DocumentSource(
            document_id="doc-4",
            filename="deck.pptx",
            content_base64=self._b64(self._pptx_bytes()),
        )
        document = parser.parse(source)
        self.assertIn("Roadmap", document.text)

    def test_pipeline_applies_chunk_quality(self) -> None:
        pipeline = IngestionPipeline()
        artifact = pipeline.ingest(
            DocumentSource(
                document_id="doc-5",
                filename="notes.txt",
                text=(
                    "This is a complete paragraph about revenue growth.\n\n"
                    "This is another paragraph with context for planning."
                ),
            )
        )
        self.assertGreaterEqual(artifact.metadata["average_chunk_quality"], 0.0)
        self.assertGreaterEqual(len(artifact.chunks), 1)
        self.assertIn("page_count", artifact.metadata)
        self.assertIn("table_count", artifact.metadata)

    @staticmethod
    def _pdf_bytes(text: str) -> bytes:
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), text)
        data = document.tobytes()
        document.close()
        return data

    @staticmethod
    def _docx_bytes() -> bytes:
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Hello world")
        document.add_paragraph("Quarterly update")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Vendor"
        table.cell(0, 1).text = "Amount"
        table.cell(1, 0).text = "ACME"
        table.cell(1, 1).text = "120"
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _xlsx_bytes() -> bytes:
        buffer = io.BytesIO()
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Quarter", "Revenue"])
        worksheet.append(["Q1", 120])
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    @staticmethod
    def _pptx_bytes() -> bytes:
        buffer = io.BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Roadmap"
        presentation.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _b64(data: bytes) -> str:
        return base64.b64encode(data).decode("ascii")


class ChunkingTests(unittest.TestCase):
    def test_structural_documents_choose_structural_strategy(self) -> None:
        engine = AntiChunkingEngine()
        document = ParsedDocument(
            document_id="doc-6",
            filename="table.csv",
            media_type="text/csv",
            parser_name="csv",
            text="a | b\n1 | 2",
            metadata={"rows": 2},
        )
        strategy, chunks = engine.chunk(document)
        self.assertEqual(strategy, "structural")
        self.assertGreaterEqual(len(chunks), 1)

    def test_quality_middleware_rechunks_low_score_chunks(self) -> None:
        middleware = ChunkQualityMiddleware(threshold=0.6)
        document = ParsedDocument(
            document_id="doc-7",
            filename="notes.txt",
            media_type="text/plain",
            parser_name="text",
            text="alpha beta gamma. delta epsilon zeta.",
        )
        original_strategy, chunks = AntiChunkingEngine().chunk(document)
        self.assertEqual(original_strategy, "semantic")
        rescored = middleware.apply(document, chunks)
        self.assertGreaterEqual(len(rescored), 1)

    def test_structural_chunks_preserve_table_provenance(self) -> None:
        artifact = IngestionPipeline().ingest(
            DocumentSource(
                document_id="doc-table",
                filename="table.csv",
                text="name,value\nrevenue,42",
            )
        )
        self.assertEqual(artifact.strategy, "structural")
        table_chunks = [chunk for chunk in artifact.chunks if chunk.metadata.get("element_type") == "table"]
        self.assertGreaterEqual(len(table_chunks), 1)
        self.assertIn("table_id", table_chunks[0].metadata)


class GraphBuilderTests(unittest.TestCase):
    def test_graph_builder_uses_networkx_payload(self) -> None:
        artifact = IngestionPipeline().ingest(
            DocumentSource(
                document_id="doc-graph",
                filename="graph.txt",
                text="Revenue increased because ACME expanded revenue operations.",
            )
        )
        payload = SimpleGraphBuilder().build([artifact])
        self.assertGreater(payload["metrics"]["node_count"], 0)
        self.assertTrue(any(edge["relation"] == "MENTIONS" for edge in payload["edges"]))


if __name__ == "__main__":
    unittest.main()
